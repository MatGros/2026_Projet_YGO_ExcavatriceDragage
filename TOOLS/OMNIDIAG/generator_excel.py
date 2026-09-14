"""
===============================================================================
📊 OMNIDIAG — Générateur d'Exports Excel (.xlsx) et CSV
===============================================================================
🎯 Rôle : Produit un classeur Excel multi-onglets mis en page de niveau manuel
   technique ainsi qu'un export CSV universel.
===============================================================================
"""

import csv
from pathlib import Path
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_to_csv(items: List[Dict[str, Any]], output_path: Path):
    """Exporte l'ensemble des diagnostics vers un fichier CSV encodé en UTF-8-SIG."""
    fields = [
        ("id", "ID"),
        ("category", "Catégorie"),
        ("organ", "Organe"),
        ("code", "Code Défaut"),
        ("text", "Message / Alarme affichée"),
        ("level", "Niveau"),
        ("cause_racine", "Cause"),
        ("action_conducteur", "Action Utilisateur"),
        ("action_maintenance", "Action Maintenance"),
        ("points_test", "Points de Test Matériels"),
        ("condition", "Condition Déclenchement (Code ST)"),
        ("source_file", "Fichier Source"),
        ("line", "Ligne"),
    ]

    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([label for _, label in fields])
        for it in items:
            row = [it.get(k, "") for k, _ in fields]
            writer.writerow(row)


def style_sheet(ws, title_text: str):
    """Applique un style professionnel aux onglets Excel."""
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=9)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Style Header
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Style Data
    for row in range(2, ws.max_row + 1):
        is_even = (row % 2 == 0)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Ajustement largeur colonnes
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            first_line = val_str.split("\n")[0]
            max_len = max(max_len, len(first_line))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_items_to_ws(ws, items: List[Dict[str, Any]]):
    """Écrit une liste d'items dans une feuille Excel."""
    headers = [
        ("id", "ID"),
        ("organ", "Organe"),
        ("code", "Code"),
        ("text", "Message / Alarme"),
        ("level", "Niveau"),
        ("cause_racine", "Cause"),
        ("action_conducteur", "Action Utilisateur"),
        ("action_maintenance", "Action Maintenance"),
        ("points_test", "Points de Test"),
        ("source_file", "Fichier"),
        ("line", "Ligne"),
    ]

    ws.append([label for _, label in headers])
    for it in items:
        row = [it.get(k, "") for k, _ in headers]
        ws.append(row)


def export_to_excel(items: List[Dict[str, Any]], output_path: Path):
    """Génère un classeur Excel structuré avec plusieurs onglets par métier."""
    wb = openpyxl.Workbook()
    # Retrait de la feuille par défaut
    wb.remove(wb.active)

    # Onglet 1 : Synthèse globale
    ws_all = wb.create_sheet(title="Toutes les Données")
    write_items_to_ws(ws_all, items)
    style_sheet(ws_all, "Toutes les Données")

    # Onglet 2 : Alarmes Carrousel
    alarms = [it for it in items if "ALARM" in it["category"]]
    ws_alm = wb.create_sheet(title="Alarmes Carrousel")
    write_items_to_ws(ws_alm, alarms)
    style_sheet(ws_alm, "Alarmes Carrousel")

    # Onglet 3 : Guidage & Actions Opérateur
    actions = [it for it in items if it["category"] == "OPERATOR_ACTION" or it["category"] == "SPECIAL_CONDITION"]
    ws_act = wb.create_sheet(title="Guidage & Actions")
    write_items_to_ws(ws_act, actions)
    style_sheet(ws_act, "Guidage & Actions")

    # Onglet 4 : Arrêt d'Urgence & Armement
    safety = [it for it in items if it["category"] in ["SECURITE_AU", "ABORT_REARMEMENT_AU"]]
    ws_safe = wb.create_sheet(title="Sécurité & AU")
    write_items_to_ws(ws_safe, safety)
    style_sheet(ws_safe, "Sécurité & AU")

    # Onglet 5 : Homing & Preflight
    homing_pref = [it for it in items if it["category"] in ["GUIDAGE_HOMING", "CHECKLIST_PREFLIGHT"]]
    ws_hom = wb.create_sheet(title="Homing & Preflight")
    write_items_to_ws(ws_hom, homing_pref)
    style_sheet(ws_hom, "Homing & Preflight")

    # Onglet 6 : Trace & Socles Défauts
    traces = [it for it in items if it["category"] in ["TRACE_BLOCAGE_TERRAIN", "SOCLE_DEFAUTS_FB"]]
    ws_trc = wb.create_sheet(title="Trace & Socles Défauts")
    write_items_to_ws(ws_trc, traces)
    style_sheet(ws_trc, "Trace & Socles Défauts")

    wb.save(output_path)
